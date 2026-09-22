"""
Pasa las planillas del gimnasio a UN solo CSV normalizado.

TODAS TIENEN LA MISMA FORMA por dentro —Nombre, Rut, Tipo de plan, Cancelado,
Fecha, Término— pero cada una empieza en una fila y una columna distintas, con
el título del convenio encima. Por eso no se leen por posición: se busca la
fila que tiene «Nombre» y «Rut», y a partir de ahí se mira por el nombre de la
columna.

Lo que sale de aquí NO se toca a mano: lo lee `php artisan datos:importar-planillas`,
que es donde se decide qué entra a la base y qué no.
"""
import csv
import re
import sys
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

# La carpeta donde están las planillas del gimnasio. Se puede pasar otra como
# segundo argumento: python scripts/planillas_a_csv.py salida.csv "D:uta"
BASE = Path(sys.argv[2] if len(sys.argv) > 2 else r'D:\Tiendas\Marketing\ProGym&Estoicos\ProGym\Progym')

# Cada planilla con el convenio al que pertenece. La maestra no lleva ninguno:
# ahí el convenio viene escrito dentro del tipo de plan, socio por socio.
PLANILLAS = [
    ('Planillas/Planilla Socios definitva 1.xlsx', None),
    ('Planillas/Convenio AIEP.xlsx', 'AIEP'),
    ('Planillas/Convenio Afusam.xlsx', 'AFUSAM'),
    ('Planillas/Convenio CDH y Taurus.xlsx', 'CDH y Taurus'),
    ('Planillas/Convenio CMPC.xlsx', 'CMPC'),
    ('Planillas/Convenio Central Park.xlsx', 'Central Park'),
    ('Planillas/Convenio Four Points y Enjoy.xlsx', 'Four Points y Enjoy'),
    ('Planillas/Convenio Hotel Muso.xlsx', 'Hotel Muso'),
    ('Planillas/Convenio Iberia.xlsx', 'Iberia'),
    ('Planillas/Convenio Inacap.xlsx', 'INACAP'),
    ('Planillas/Convenio Los Angeles VIP y Divas Eley Urbano Estilo.xlsx', 'Los Ángeles VIP y Divas'),
    ('Planillas/Convenio Masisa.xlsx', 'Masisa'),
    ('Planillas/Convenio Promasa.xlsx', 'Promasa'),
    ('Planillas/Convenio Ripley.xlsx', 'Ripley'),
    ('Planillas/Convenio Saint George.xlsx', 'Saint George'),
    ('Planillas/Convenio San Rafael.xlsx', 'San Rafael'),
    ('Planillas/Convenio SoloDivas.xlsx', 'Solo Divas'),
    ('Planillas/Convenio Sto Tomas.xlsx', 'Santo Tomás'),
    ('Planillas/Convenio UCSC.xlsx', 'UCSC'),
    ('Planillas/Convenio UDEC.xlsx', 'Universidad de Concepción'),
    ('Planillas/Convenio Virginio Gomez.xlsx', 'IP Virginio Gómez'),
    ('Planillas/Planilla Hispanoamericano.xlsx', 'Colegio Hispanoamericano'),
    ('Planillas/Planilla Deluxe.xlsx', 'Deluxe'),
    ('Planillas/Planilla Millalen.xlsx', 'Millalén'),
    ('Planillas/Planilla Nestle.xlsx', 'Nestlé'),
    ('Planillas/Socios Estilo Latino.xlsx', 'Estilo Latino'),
    ('Planillas/El Molino.xlsx', 'El Molino'),
    ('Planillas/Shivi.xlsx', 'Shivi'),
    ('Planillas/Ceala.xlsx', 'Ceala'),
    ('Planillas/Jumbo,Paris, Cencosud,Easy,Falabella.xlsx', 'Cencosud (Jumbo, Paris, Easy)'),
    ('Planillas/Planilla Jormat 2024.xlsx', 'Jormat'),
    ('Planillas/Planilla TRABAJADORES JORMAT.xlsx', 'Jormat'),
]

# Lo que NO entra, y por qué:
#  · PlanillSocios.xlsx, Planilla Socios definitva.xlsx, PlanillSocios (1).xlsx
#    y Convenio Jumbo.xlsx  → son de 2022 y la maestra ya los trae.
#  · Los archivos -DESKTOP-T2LIBA4 → copias del mismo archivo.
#  · Nomina Promasa S.A_.xlsx → la nómina de la empresa, no socios del gimnasio.
#  · Deudas Junio 2026, Cuentas Plusport, Precios productos, Inventario Suples,
#    ingresosgastos → suplementos y cuentas personales, no membresías.

COLUMNAS = {
    'nombre': 'nombre',
    'nombres': 'nombre',
    'rut': 'rut',
    'run': 'rut',
    'tipo de plan': 'plan',
    'plan': 'plan',
    'plan anual': 'plan_anual',
    'cancelado': 'monto',
    'fecha': 'inicio',
    'termino': 'termino',
    'término': 'termino',
    'cel': 'celular',
    'celular': 'celular',
    'telefono': 'celular',
    'teléfono': 'celular',
}


def normalizar(cabecera) -> str:
    texto = re.sub(r'\s+', ' ', str(cabecera or '')).strip().lower()
    return COLUMNAS.get(texto, '')


def texto(valor) -> str:
    if valor is None:
        return ''
    if isinstance(valor, datetime):
        return valor.strftime('%Y-%m-%d')
    return re.sub(r'\s+', ' ', str(valor)).strip()


def leer(ruta: Path):
    """Devuelve las filas de una planilla como diccionarios."""
    wb = load_workbook(ruta, data_only=True)
    hoja = wb.worksheets[0]
    mapa, encabezado = None, 0

    for i, fila in enumerate(hoja.iter_rows(values_only=True), start=1):
        etiquetas = [normalizar(c) for c in fila]

        if mapa is None:
            # La fila de los títulos es la que trae a la vez el nombre y el rut.
            if 'nombre' in etiquetas and 'rut' in etiquetas:
                mapa = {etiqueta: j for j, etiqueta in enumerate(etiquetas) if etiqueta}
                encabezado = i
            continue

        datos = {campo: texto(fila[j]) if j < len(fila) else '' for campo, j in mapa.items()}

        if not datos.get('nombre'):
            continue

        # Jormat: el plan va en la cabecera («PLAN ANUAL») y la celda lleva el precio.
        if not datos.get('plan') and datos.get('plan_anual'):
            datos['plan'] = 'anual'
            datos['monto'] = datos['plan_anual']

        datos.pop('plan_anual', None)
        yield datos

    if mapa is None:
        print(f'   !! sin cabecera reconocible: {ruta.name}', file=sys.stderr)


salida = Path(sys.argv[1] if len(sys.argv) > 1 else 'planillas.csv')
total = 0

with salida.open('w', encoding='utf-8', newline='') as f:
    w = csv.writer(f)
    w.writerow(['planilla', 'convenio', 'nombre', 'rut', 'plan', 'monto', 'inicio', 'termino', 'celular'])

    for relativa, convenio in PLANILLAS:
        ruta = BASE / relativa
        if not ruta.exists():
            print(f'   !! no está: {relativa}', file=sys.stderr)
            continue

        cuantas = 0
        for fila in leer(ruta):
            w.writerow([
                ruta.name,
                convenio or '',
                fila.get('nombre', ''),
                fila.get('rut', ''),
                fila.get('plan', ''),
                fila.get('monto', ''),
                fila.get('inicio', ''),
                fila.get('termino', ''),
                fila.get('celular', ''),
            ])
            cuantas += 1

        total += cuantas
        print(f'{cuantas:5}  {ruta.name}  →  {convenio or "(la maestra)"}')

print(f'\n{total} filas en {salida}')
